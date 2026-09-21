import requests
import json
import os
import time
import uuid

from datetime import datetime
from fastapi import FastAPI
from pydantic import BaseModel
from requests_negotiate_sspi import HttpNegotiateAuth
from io import BytesIO
from docx import Document
from pypdf import PdfReader
from openai import OpenAI
from fastapi.middleware.cors import CORSMiddleware
from fastapi import FastAPI, HTTPException


app = FastAPI()
app.add_middleware(
    CORSMiddleware,
    allow_origins=[
        "http://calsawsproject",
    ],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


class AskRequest(BaseModel):
    question: str


LOG_DIR = "logs"
LOG_FILE = os.path.join(LOG_DIR, "chatbot.jsonl")

os.makedirs(LOG_DIR, exist_ok=True)


def write_chat_log(entry):
    with open(LOG_FILE, "a", encoding="utf-8") as f:
        f.write(
            json.dumps(entry, ensure_ascii=False) + "\n"
        )
        

def make_search_query(question):
    client = OpenAI()

    response = client.responses.create(
        model="gpt-5.6",
        instructions=(
            "Extract only the most important keywords and named terms "
            "from the user's question for use in SharePoint Search. "
            "Do not include words describing the user's intent, such as "
            "'what', 'how', 'why', 'topics', 'information', 'meeting', "
            "'documents', or 'details'. "
            "Return 2 to 4 search terms only, separated by spaces. "
            "Do not use quotation marks."
        ),
        input=question
    )

    return response.output_text.strip()
    
    
def search_sharepoint(query):
    url = "http://calsawsproject/_api/search/query"

    params = {
        "querytext": f"'{query}'"
    }

    response = requests.get(
        url,
        params=params,
        auth=HttpNegotiateAuth(),
        headers={
            "Accept": "application/json;odata=verbose"
        }
    )

    response.raise_for_status()

    data = response.json()

    rows = data["d"]["query"]["PrimaryQueryResult"][
        "RelevantResults"
    ]["Table"]["Rows"]["results"]

    results = []

    for row in rows:
        cells = {
            cell["Key"]: cell["Value"]
            for cell in row["Cells"]["results"]
        }

        results.append({
            "title": cells.get("Title"),
            "path": cells.get("Path"),
            "file_type": cells.get("FileType"),
            "rank": cells.get("Rank"),
            "snippet": cells.get("HitHighlightedSummary")
        })

    return results


def get_docx_text(url):
    response = requests.get(
        url,
        auth=HttpNegotiateAuth()
    )

    response.raise_for_status()

    document = Document(BytesIO(response.content))

    parts = []

    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            parts.append(paragraph.text)

    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            parts.append(" | ".join(cells))

    return "\n".join(parts)


def get_pdf_text(url):
    response = requests.get(
        url,
        auth=HttpNegotiateAuth()
    )

    response.raise_for_status()

    reader = PdfReader(BytesIO(response.content))

    parts = []

    for page in reader.pages:
        text = page.extract_text()

        if text and text.strip():
            parts.append(text)

    return "\n".join(parts)
    
    
@app.get("/")
def root():
    return {"status": "SharePoint Bot API is running"}
  
    
@app.post("/ask")
def ask(request: AskRequest):

    start_time = time.perf_counter()
    
    request_id = str(uuid.uuid4())
    print("REQUEST ID:", request_id)

    # Keep track of the processing stage for error diagnostics
    stage = "start"
    search_query = None

    try:
        # ---------------------------------------------------------
        # 1. Generate SharePoint search query
        # ---------------------------------------------------------
        stage = "make_search_query"

        search_query = make_search_query(request.question)
        print("SEARCH QUERY:", search_query)

        # ---------------------------------------------------------
        # 2. Search SharePoint
        # ---------------------------------------------------------
        stage = "sharepoint_search"

        results = search_sharepoint(search_query)

        # ---------------------------------------------------------
        # 3. Build context from SharePoint search results
        # ---------------------------------------------------------
        stage = "build_context"

        context_parts = []

        for result in results:
            context_parts.append(
                f"""
Document: {result["title"]}
Path: {result["path"]}
File type: {result["file_type"]}
Rank: {result["rank"]}
Snippet: {result["snippet"]}
"""
            )

        context = "\n\n---\n\n".join(context_parts)

        # ---------------------------------------------------------
        # 4. Generate answer with OpenAI
        # ---------------------------------------------------------
        stage = "openai_answer"
        client = OpenAI()

        response = client.responses.create(
            model="gpt-5.6",
            instructions=(
                "Answer the user's question using only the information "
                "provided in the SharePoint search results. "
                "If the search results do not contain enough information, "
                "say that the information was not found."
            ),
            input=f"""
User question:
{request.question}

SharePoint search results:
{context}
"""
        )

        usage = response.usage

        duration_ms = round(
            (time.perf_counter() - start_time) * 1000
        )

        print("OPENAI USAGE:")
        print("  Input tokens :", usage.input_tokens)
        print(
            "  Cached tokens:",
            usage.input_tokens_details.cached_tokens
        )
        print("  Output tokens:", usage.output_tokens)
        print("  Total tokens :", usage.total_tokens)
        print("  Duration     :", duration_ms, "ms")

        # ---------------------------------------------------------
        # 5. Log successful request
        # ---------------------------------------------------------
        log_entry = {
            "timestamp": datetime.now().astimezone().isoformat(),
            "status": "success",
            "request_id": request_id,
            "question": request.question,
            "search_query": search_query,
            "answer": response.output_text,
            "usage": {
                "input_tokens": usage.input_tokens,
                "cached_tokens":
                    usage.input_tokens_details.cached_tokens,
                "output_tokens": usage.output_tokens,
                "total_tokens": usage.total_tokens
            },
            "duration_ms": duration_ms,
            "results": [
                {
                    "title": result.get("title"),
                    "path": result.get("path"),
                    "file_type": result.get("file_type"),
                    "rank": result.get("rank"),
                    "snippet": result.get("snippet")
                }
                for result in results
            ]
        }

        write_chat_log(log_entry)

        return {
            "request_id": request_id,
            "question": request.question,
            "search_query": search_query,
            "answer": response.output_text,
            "results": results
        }

    except Exception as e:

        duration_ms = round(
            (time.perf_counter() - start_time) * 1000
        )

        print("ERROR:")
        print("  Stage    :", stage)
        print("  Type     :", type(e).__name__)
        print("  Message  :", str(e))
        print("  Duration :", duration_ms, "ms")

        # ---------------------------------------------------------
        # Log failed request
        # ---------------------------------------------------------
        error_entry = {
            "timestamp": datetime.now().astimezone().isoformat(),
            "status": "error",
            "request_id": request_id,
            "question": request.question,
            "search_query": search_query,
            "stage": stage,
            "error_type": type(e).__name__,
            "error": str(e),
            "duration_ms": duration_ms
        }

        write_chat_log(error_entry)

        raise
        
        
class FeedbackRequest(BaseModel):
    request_id: str
    feedback: str


@app.post("/feedback")
def feedback(request: FeedbackRequest):

    if request.feedback not in ("up", "down"):
        raise HTTPException(
            status_code=400,
            detail="Invalid feedback value"
        )

    log_entry = {
        "timestamp": datetime.now().astimezone().isoformat(),
        "status": "feedback",
        "request_id": request.request_id,
        "feedback": request.feedback
    }

    write_chat_log(log_entry)

    return {
        "status": "ok"
    }  
    
    
@app.get("/stats")
def get_stats():

    request_count = 0

    if os.path.exists(LOG_FILE):
        with open(LOG_FILE, "r", encoding="utf-8") as f:
            for line in f:
                try:
                    entry = json.loads(line)

                    if entry.get("status") == "success":
                        request_count += 1

                except json.JSONDecodeError:
                    continue

    return {
        "requests": request_count
    }